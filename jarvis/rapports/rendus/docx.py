"""Rendu Word (python-docx), a partir du meme modele de document que le HTML.

Construit directement en Word (titres, legendes, tableaux stylés), pas par
conversion du HTML: le document reste editable proprement par l'utilisateur.
"""
import io

from ..document import Encadre, Figure, LIBELLES_ORIGINE, Liste, Paragraphe, Tableau
from ..faits import LIBELLES_STATUT

ACCENT = (0x0B, 0x6E, 0x8A)
COULEURS_STATUT = {"observe": (0x1F, 0x7A, 0x4D), "correle": (0x8A, 0x5A, 0x00),
                   "projete": (0x6B, 0x3F, 0xA0), "methode": (0x4A, 0x58, 0x63)}
COULEURS_ENCADRE = {"limite": (0x9A, 0x34, 0x12), "officiel": (0x0B, 0x4F, 0x8A),
                    "avertissement": (0x8A, 0x5A, 0x00), "hypothese": (0x4A, 0x58, 0x63),
                    "information": (0x4A, 0x58, 0x63)}


def _couleur(run, rgb):
    from docx.shared import RGBColor
    run.font.color.rgb = RGBColor(*rgb)


def _legende(doc, gras, texte, meta):
    from docx.shared import Pt
    p = doc.add_paragraph()
    r = p.add_run(gras)
    r.bold = True
    r.font.size = Pt(8.5)
    r = p.add_run(" " + texte)
    r.font.size = Pt(8.5)
    p2 = doc.add_paragraph()
    r = p2.add_run(meta)
    r.font.size = Pt(8)
    r.italic = True
    _couleur(r, (0x4A, 0x58, 0x63))


def _meta_visuel(b):
    texte = "Unité : %s · Période : %s · Source : %s" % (b.unite, b.periode, b.source)
    if b.statut:
        texte += " · Statut : %s" % LIBELLES_STATUT[b.statut]
    if isinstance(b, Figure):
        texte += " · " + LIBELLES_ORIGINE[b.origine]
    return texte


def en_docx(rapport) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    m = rapport.meta
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(1.8)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    marque = doc.add_paragraph()
    r = marque.add_run("ClimatSen · RAPPORT GÉNÉRÉ PAR IRIS")
    r.bold = True
    r.font.size = Pt(8.5)
    _couleur(r, ACCENT)
    titre = doc.add_heading(m.titre, level=0)
    titre.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st = doc.add_paragraph()
    r = st.add_run(m.sous_titre)
    r.font.size = Pt(12)
    _couleur(r, (0x4A, 0x58, 0x63))
    infos = doc.add_table(rows=0, cols=2)
    infos.style = "Light List"
    for cle, val in (("Zone", m.zone), ("Période", m.periode), ("Public", m.public),
                     ("Généré le", m.genere_le),
                     ("Version des données", "%s (code %s)" % (m.version_donnees, m.commit)),
                     ("Référence", m.id)):
        cellules = infos.add_row().cells
        cellules[0].text = cle
        cellules[1].text = val
        for c in cellules:
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(9)

    for i, s in enumerate(rapport.dans_l_ordre(), 1):
        doc.add_heading("%d. %s" % (i, s.titre), level=1)
        for b in s.blocs:
            if isinstance(b, Paragraphe):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if b.statut:
                    r = p.add_run("[%s] " % LIBELLES_STATUT[b.statut].upper())
                    r.bold = True
                    r.font.size = Pt(8)
                    _couleur(r, COULEURS_STATUT[b.statut])
                p.add_run(b.texte)
            elif isinstance(b, Liste):
                for e in b.elements:
                    doc.add_paragraph(e, style="List Number" if b.ordonnee else "List Bullet")
            elif isinstance(b, Encadre):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                if b.titre:
                    r = p.add_run(b.titre + " — ")
                    r.bold = True
                    _couleur(r, COULEURS_ENCADRE[b.genre])
                r = p.add_run(b.texte)
                r.font.size = Pt(9.5)
            elif isinstance(b, Figure):
                doc.add_picture(io.BytesIO(b.png), width=Cm(16.5))
                _legende(doc, "Figure %d — %s." % (b.numero, b.titre), b.legende, _meta_visuel(b))
            elif isinstance(b, Tableau):
                cap = doc.add_paragraph()
                r = cap.add_run("Tableau %d — %s" % (b.numero, b.titre))
                r.bold = True
                r.font.size = Pt(9)
                t = doc.add_table(rows=1, cols=len(b.colonnes))
                t.style = "Light Grid Accent 1"
                for c, nom in zip(t.rows[0].cells, b.colonnes):
                    c.text = str(nom)
                for ligne in b.lignes:
                    cellules = t.add_row().cells
                    for c, val in zip(cellules, ligne):
                        c.text = str(val)
                for row in t.rows:
                    for c in row.cells:
                        for par in c.paragraphs:
                            for run in par.runs:
                                run.font.size = Pt(8)
                _legende(doc, "", b.legende, _meta_visuel(b))

    pied = doc.sections[0].footer.paragraphs[0]
    r = pied.add_run("ClimatSen · %s · généré le %s · données %s · complète les alertes "
                     "officielles de l'ANACIM sans les remplacer" % (
                         m.id, m.genere_le, m.version_donnees))
    r.font.size = Pt(7)
    _couleur(r, (0x4A, 0x58, 0x63))
    props = doc.core_properties
    props.title = "%s — %s" % (m.titre, m.sous_titre)
    props.author = "ClimatSen (Iris)"
    props.subject = m.id
    props.keywords = "version des données %s" % m.version_donnees
    sortie = io.BytesIO()
    doc.save(sortie)
    return sortie.getvalue()
