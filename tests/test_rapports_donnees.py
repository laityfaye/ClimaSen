"""Rapports sur les VRAIES donnees: chiffres, fidelite, rendus.

Chaque type de rapport est produit une fois (gabarit, sans reseau). On
verifie que ses faits sont ceux des fichiers sources, que tout nombre de la
prose vient d'un fait ou d'un texte ecrit par le code, et que les trois
rendus portent la meme structure, la version des donnees et chaque visuel.
"""
import asyncio
import datetime as dt
import io
import json
import re

import pandas as pd
import pytest

from jarvis.rapports import gazetteer, manifest, redaction, service, verification
from jarvis.rapports import spec as S
from jarvis.rapports import textes_fixes as tf
from jarvis.rapports.document import CLES_SECTIONS, Liste, Paragraphe
from jarvis.rapports.rendus import en_docx, en_html
from jarvis.tools import dataset

AUJ = dt.date(2026, 10, 7)
CAS = {
    "vul_pikine": {"type": "vulnerabilite", "zone": "Pikine"},
    "vul_mbao": {"type": "vulnerabilite", "zone": "Mbao"},
    "vul_national": {"type": "vulnerabilite", "zone_code": "SN"},
    "vul_region": {"type": "vulnerabilite", "zone": "region de Kolda"},
    "hist_kolda": {"type": "historique", "zone_code": "SN07"},
    "hist_pikine": {"type": "historique", "zone": "Pikine"},
    "hist_phase": {"type": "historique", "phase": "Phase_2_pleine", "year_min": 2000},
    "tele": {"type": "teleconnexions"},
    "tele_p3": {"type": "teleconnexions", "phase": "Phase_3_fin", "metric": "n_events"},
    "veille_2027": {"type": "veille", "horizon": "prochaine", "zone": "Pikine"},
    "veille_2023": {"type": "veille", "season": 2023},
}


@pytest.fixture(scope="module")
def produits():
    from veille import production
    g = gazetteer.charger()
    sortie = {}
    for nom, params in CAS.items():
        r = S.preparer(params, g, production.bulletins_disponibles(), AUJ)
        assert r["statut"] == "pret", (nom, r)
        sortie[nom] = asyncio.run(service.produire("R%s" % nom.replace("_", ""), r["spec"],
                                                   client=None, gazetteer=g, avec_pdf=False))
    return sortie


def _fait(produits, nom, ident):
    return produits[nom][1].registre[ident]


# =============================================================================
# Les chiffres sont ceux des fichiers
# =============================================================================
def test_vulnerabilite_pikine_egale_au_csv(produits):
    csv = pd.read_csv("outputs/vulnerabilite/indice_risque_departements.csv")
    ligne = csv[csv["pcode"] == "SN0103"].iloc[0]
    assert _fait(produits, "vul_pikine", "zone_rang").valeur == ligne["rang"]
    assert _fait(produits, "vul_pikine", "zone_indice").valeur == pytest.approx(
        ligne["indice_risque"], abs=1e-3)
    assert _fait(produits, "vul_pikine", "zone_pop2023").valeur == ligne["population_2023"]
    assert _fait(produits, "vul_pikine", "zone_sur").valeur == len(csv)
    assert "biais_dakar" in produits["vul_pikine"][1].limites


def test_commune_decrite_par_son_arrondissement(produits):
    col = produits["vul_mbao"][1]
    com = pd.read_csv("data/processed/correspondance_communes_arrondissements.csv")
    pop = com[com["COMMUNE"] == "MBAO"]["population_2023"].iloc[0]
    assert col.registre["commune_pop2023"].valeur == pop
    assert col.registre["zone_nom"].valeur == "Thiaroye"
    assert "echelle_arrondissement" in col.limites


def test_classement_national(produits):
    csv = pd.read_csv("outputs/vulnerabilite/indice_risque_departements.csv")
    premier = csv.sort_values("indice_risque", ascending=False).iloc[0]
    assert _fait(produits, "vul_national", "top1_nom").valeur == premier["departement"]


def test_historique_kolda_egal_au_catalogue(produits):
    ev = dataset.get("events")
    sel = ev[ev["max_intensity_region"] == "Kolda"]
    assert _fait(produits, "hist_kolda", "n_zone").valeur == len(sel)
    assert _fait(produits, "hist_kolda", "n_national").valeur == len(ev)
    assert _fait(produits, "hist_kolda", "record_mm").valeur == pytest.approx(sel["max_precip"].max())
    for ph in ("Phase_1_debut", "Phase_2_pleine", "Phase_3_fin"):
        assert _fait(produits, "hist_kolda", "n_" + ph).valeur == (sel["phase"] == ph).sum()


def test_historique_filtres_phase_et_periode(produits):
    ev = dataset.get("events")
    sel = ev[(ev["phase"] == "Phase_2_pleine") & (ev["year"] >= 2000)]
    assert _fait(produits, "hist_phase", "n_national").valeur == len(sel)
    assert "periode_reduite" in produits["hist_phase"][1].limites


def test_historique_sans_evenement_dit_pourquoi(produits):
    col = produits["hist_pikine"][1]
    assert col.registre["n_zone"].valeur == 0
    assert "local_jours_2sigma" in col.registre


def test_teleconnexions_egales_au_script_04(produits):
    df = pd.read_csv("outputs/teleconnections/correlations_Phase_2_pleine.csv")
    df = df[df["metric"] == "max_precip"]
    reg = produits["tele"][1].registre
    ligne = df[(df["index"] == reg["p2_top_indice"].valeur)
               & (df["lag_months"] == reg["p2_top_lag"].valeur)].iloc[0]
    assert reg["p2_top_r"].valeur == pytest.approx(ligne["pearson_r"], abs=1e-3)
    assert reg["p2_n_tests"].valeur == len(df)
    assert reg["p2_n_sig"].valeur == (df["pearson_p_neff"] < 0.05).sum()
    assert reg["p2_attendues"].valeur == pytest.approx(len(df) * 0.05)
    # Intervalle de confiance: contient r, et exclut 0 pour une correlation significative.
    assert reg["p2_top_ic_bas"].valeur < reg["p2_top_r"].valeur < reg["p2_top_ic_haut"].valeur
    assert reg["p2_top_ic_haut"].valeur < 0
    assert "multiplicite_tests" in produits["tele"][1].limites


def test_teleconnexions_une_seule_phase(produits):
    reg = produits["tele_p3"][1].registre
    assert "p3_n_tests" in reg and "p2_n_tests" not in reg


def test_ic_fisher():
    from jarvis.rapports.collecteurs.teleconnexions import ic_fisher
    bas, haut = ic_fisher(0.0, 41)
    assert bas == pytest.approx(-haut) and 0.29 < haut < 0.32
    assert ic_fisher(0.5, 3) == (None, None)


def test_veille_sans_prevision_n_annonce_aucun_niveau(produits):
    rapport, col = produits["veille_2027"][:2]
    assert col.registre["saison"].valeur == 2027
    assert "niveau" not in col.registre and "proba" not in col.registre
    textes = " ".join(rapport.textes())
    assert tf.MENTION_ANACIM in textes


def test_veille_saison_passee_probabilite_sans_niveau(produits):
    reg = produits["veille_2023"][1].registre
    assert reg["proba"].statut == "projete"
    assert "niveau" not in reg            # mode "probabilite": competence non demontree
    assert "c3s_auc" in reg and "verif_extreme" in reg


def test_bulletin_d_une_saison_passee_dit_qu_il_est_reconstitue(produits):
    """Les bulletins 1998-2023 ont ete calcules apres coup: le rapport ne doit
    pas dire qu'ils font le point "a la date du" jour de calcul."""
    passe = " ".join(produits["veille_2023"][0].textes())
    assert "reconstitution a posteriori" in passe and "à la date du" not in passe
    avenir = " ".join(produits["veille_2027"][0].textes())
    assert "à la date du" in avenir


# =============================================================================
# Fidelite: aucun nombre de la prose ne vient d'ailleurs
# =============================================================================
@pytest.mark.parametrize("nom", list(CAS))
def test_chaque_nombre_de_la_prose_est_trace(produits, nom):
    rapport, col, journal, _ = produits[nom]
    reg = col.registre
    permis = verification.autorises(reg, redaction.textes_du_code(col, _spec_de(nom)))
    permis.add(str(len(reg)))            # "le registre des N faits" (annexe)
    textes_code = (list(tf.LIMITES.values()) + list(tf.METHODOLOGIES.values())
                   + [tf.MENTION_ANACIM, tf.AVERTISSEMENT_CORRELATION, tf.LEGENDE_STATUTS,
                      tf.AVERTISSEMENT_PROJECTION, tf.AVERTISSEMENT_INDICE]
                   + list(tf.SOURCES.values()) + col.limites_specifiques)
    for t in textes_code:
        permis.update(verification.nombres(t))
    permis.update(verification.nombres(rapport.meta.version_donnees + " " + rapport.meta.id
                                       + " " + rapport.meta.genere_le + " " + rapport.meta.commit))
    for section in rapport.dans_l_ordre():
        for b in section.blocs:
            textes = [b.texte] if isinstance(b, Paragraphe) else (
                b.elements if isinstance(b, Liste) else [])
            for t in textes:
                hors = [n for n in verification.nombres(t) if n not in permis]
                assert not hors, (nom, section.cle, t, hors)


def _spec_de(nom):
    from veille import production
    return S.preparer(CAS[nom], gazetteer.charger(), production.bulletins_disponibles(),
                      AUJ)["spec"]


@pytest.mark.parametrize("nom", list(CAS))
def test_structure_et_metadonnees(produits, nom):
    rapport = produits[nom][0]
    assert [s.cle for s in rapport.dans_l_ordre()] == list(CLES_SECTIONS)
    assert all(s.blocs for s in rapport.dans_l_ordre())
    assert rapport.meta.version_donnees == manifest.lire()["version_donnees"]
    for f in rapport.figures() + rapport.tableaux():
        assert f.titre and f.legende and f.unite and f.periode and f.source
    recos = [b for b in rapport.sections["conclusions"].blocs if isinstance(b, Liste)]
    assert recos and recos[0].ordonnee and recos[0].elements


@pytest.mark.parametrize("nom", list(CAS))
def test_aucun_renvoi_oublie(produits, nom):
    assert "{{" not in " ".join(produits[nom][0].textes())


# =============================================================================
# Rendus
# =============================================================================
@pytest.mark.parametrize("nom", ["vul_pikine", "tele", "veille_2027"])
def test_html_complet(produits, nom):
    rapport = produits[nom][0]
    html = en_html(rapport)
    for s in rapport.dans_l_ordre():
        assert s.titre.replace("'", "&#39;") in html
    assert rapport.meta.version_donnees in html
    assert html.count("data:image/png;base64,") == len(rapport.figures())
    assert html.count("<table>") == len(rapport.tableaux())
    assert "Unité :" in html and "Source :" in html and "Période :" in html


def test_html_echappe_les_textes():
    from jarvis.rapports.document import Meta, Rapport
    r = Rapport(Meta(id="R1", type="historique", titre="<script>alert(1)</script>",
                     sous_titre="S", genere_le="g", version_donnees="D", commit="c",
                     redaction="gabarit", modele="aucun", public="p", zone="z", periode="p"))
    for cle in CLES_SECTIONS:
        r.ajouter(cle, Paragraphe("<img src=x onerror=alert(1)>"))
    html = en_html(r.finaliser())
    assert "<script>alert" not in html and "<img src=x" not in html


@pytest.mark.parametrize("nom", ["vul_pikine", "hist_kolda"])
def test_word_complet(produits, nom):
    from docx import Document
    rapport = produits[nom][0]
    d = Document(io.BytesIO(en_docx(rapport)))
    titres = [p.text for p in d.paragraphs if p.style.name.startswith("Heading 1")]
    assert len(titres) == len(CLES_SECTIONS)
    assert len(d.inline_shapes) == len(rapport.figures())
    assert len(d.tables) == len(rapport.tableaux()) + 1      # + bloc d'informations
    assert rapport.meta.version_donnees in d.sections[0].footer.paragraphs[0].text
    assert d.core_properties.subject == rapport.meta.id


def test_pdf(produits):
    pytest.importorskip("pypdf")
    from pypdf import PdfReader
    from jarvis.rapports.rendus import ExportIndisponible, en_pdf
    rapport = produits["veille_2027"][0]
    try:
        pdf = en_pdf(en_html(rapport), rapport.meta)
    except ExportIndisponible as exc:
        pytest.skip("Chromium indisponible: %s" % exc)
    lecteur = PdfReader(io.BytesIO(pdf))
    texte = " ".join(p.extract_text() for p in lecteur.pages)
    assert len(lecteur.pages) >= 2
    assert rapport.meta.version_donnees in texte
    assert "ANACIM" in texte and "page 1" in texte
    assert re.search(r"Résumé\s+exécutif", texte)


def test_exporter_sans_pdf_renvoie_html_et_word(produits):
    sorties = service.exporter(produits["tele"][0], avec_pdf=False)
    assert sorties["html"].startswith(b"<!doctype html>")
    assert sorties["docx"][:2] == b"PK" and sorties["pdf"] is None


def test_faits_exportes_en_csv(produits):
    csv = produits["vul_pikine"][3]["faits_csv"].decode("utf-8-sig")
    assert csv.startswith("id,libelle,valeur") and "zone_indice" in csv


# =============================================================================
# Redaction IA simulee de bout en bout
# =============================================================================
class RedacteurRecopieur:
    """Renvoie le gabarit tel quel: une redaction 'parfaite' du point de vue
    des chiffres. Verifie le chemin IA complet sans reseau."""
    async def generer_json(self, system, messages, schema, profile="public", **kw):
        charge = json.loads(messages[0]["content"])
        red = dict(charge["gabarit"])
        red["recommandations"] = charge["recommandations_gabarit"]
        return {"text": json.dumps(red), "usage": {}, "model": "recopieur"}


def test_chemin_ia_complet():
    from veille import production
    g = gazetteer.charger()
    sp = S.preparer(CAS["vul_pikine"], g, production.bulletins_disponibles(), AUJ)["spec"]
    rapport, col, journal, sorties = asyncio.run(service.produire(
        "RIA", sp, client=RedacteurRecopieur(), gazetteer=g, avec_pdf=False))
    assert journal["mode"] == "ia" and rapport.meta.redaction == "ia"
    assert "recopieur" in " ".join(rapport.textes())
