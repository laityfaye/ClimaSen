# -*- coding: utf-8 -*-
"""Versions diffusables d'un bulletin, pour les acteurs operationnels.

Trois formats, composes par le CODE a partir du JSON du bulletin (aucun
modele de langage: un chiffre ne peut pas etre deforme en route):
  - message court (SMS / WhatsApp, 320 caracteres au plus);
  - resume d'une page (Markdown);
  - document Word formel (.docx), avec la fiabilite du systeme.

L'ecriture des fichiers passe par le protocole d'approbation de Jarvis
(jarvis/actions.py, type "veille_diffusion"): Jarvis propose, l'utilisateur
approuve, le serveur ecrit dans outputs/veille/diffusion/.
"""
from . import DOSSIER_SORTIE

DOSSIER = DOSSIER_SORTIE / "diffusion"
MAX_SMS = 320

CONSEILS = {
    "faible": "Maintenir la vigilance habituelle : des pluies intenses locales restent possibles.",
    "normal": "Vigilance habituelle de saison : vérifier le curage des caniveaux et les plans de contingence.",
    "eleve": "Renforcer la préparation : curage des ouvrages, prépositionnement, information des quartiers exposés.",
    "tres_eleve": "Préparation renforcée recommandée dès mai-juin : plans de contingence, zones inondables, stocks.",
    "indetermine": "Aucune prévision officielle disponible : suivre les bulletins de l'ANACIM.",
}


def _pct(p):
    return "n/d" if p is None else "%d %%" % round(100 * p)


def sms(b):
    """Message court (<= 320 caracteres)."""
    n = b["niveau_risque"]
    texte = ("CLIMAT-SEN veille %d : risque d'année extrême %s" % (b["annee"], n["libelle"].upper()))
    if n["probabilite_annee_extreme"] is not None:
        texte += " (%s, réf. 33 %%, confiance %s)" % (_pct(n["probabilite_annee_extreme"]),
                                                       n["confiance"])
    texte += ". " + CONSEILS[n["code"]] + " Alertes officielles : ANACIM."
    if len(texte) > MAX_SMS:
        texte = texte[:MAX_SMS - 1] + "…"
    return texte


def _fiabilite():
    from . import fiabilite
    c = fiabilite.carnet()
    if not c.get("disponible"):
        return None
    r = c["niveau_de_risque"]
    return ("Sur %d saisons testées (%d-%d), le système aurait détecté %d saisons extrêmes, "
            "en aurait manqué %d et émis %d fausses alertes." % (
                c["n_saisons"], c["periode"][0], c["periode"][1], r["comptes"]["détection"],
                r["comptes"]["manquée"], r["comptes"]["fausse alerte"]))


def resume(b):
    """Resume d'une page (Markdown)."""
    n = b["niveau_risque"]
    lignes = [
        "# Veille pré-saison %d — risque d'année extrême : %s" % (b["annee"], n["libelle"]),
        "",
        "**Émis le %s par CLIMAT-SEN.**" % b["emis_le"],
        "",
    ]
    if n["probabilite_annee_extreme"] is not None:
        lignes += ["- Probabilité d'une année extrême : **%s** (référence : 33 %%)" %
                   _pct(n["probabilite_annee_extreme"]),
                   "- Source : prévision saisonnière Copernicus C3S calibrée sur CHIRPS",
                   "- Confiance : **%s**" % n["confiance"]]
    else:
        lignes.append("- Niveau non déterminé : prévision officielle indisponible")
    lignes += ["", "## Recommandation", "", CONSEILS[n["code"]], "", "## Synthèse", "",
               b["synthese"], ""]
    fiab = _fiabilite()
    if fiab:
        lignes += ["## Fiabilité", "", fiab, ""]
    lignes += ["---", "", "*Une année extrême est une saison dont l'étendue cumulée des pluies "
               "extrêmes est dans le tiers supérieur (1999, 2005, 2010, 2012, 2020, 2022...). "
               "Ce bulletin ne remplace pas les prévisions et alertes de l'ANACIM.*"]
    return "\n".join(lignes) + "\n"


def docx(b, chemin):
    """Document Word formel."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    n = b["niveau_risque"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.add_heading("Bulletin de veille pré-saison — saison des pluies %d" % b["annee"], 0)
    doc.add_paragraph("CLIMAT-SEN · émis le %s · statut : %s" % (b["emis_le"], b["statut"]))

    doc.add_heading("Niveau de risque d'année extrême", 1)
    p = doc.add_paragraph()
    run = p.add_run(n["libelle"].upper())
    run.bold = True
    run.font.size = Pt(20)
    couleur = n["couleur"].lstrip("#")
    run.font.color.rgb = RGBColor(int(couleur[0:2], 16), int(couleur[2:4], 16), int(couleur[4:6], 16))
    if n["probabilite_annee_extreme"] is not None:
        doc.add_paragraph("Probabilité : %s (référence climatologique : 33 %%). Confiance : %s."
                          % (_pct(n["probabilite_annee_extreme"]), n["confiance"]))

    doc.add_heading("Recommandation", 1)
    doc.add_paragraph(CONSEILS[n["code"]])

    doc.add_heading("Synthèse", 1)
    doc.add_paragraph(b["synthese"])

    p = b.get("projection") or {}
    if p.get("analogues"):
        doc.add_heading("Saisons passées les plus ressemblantes", 1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for i, titre in enumerate(("Année", "Saison", "Inondations documentées")):
            table.rows[0].cells[i].text = titre
        for a in p["analogues"]:
            cells = table.add_row().cells
            cells[0].text = str(a["annee"])
            cells[1].text = "extrême" if a["extreme"] else "normale"
            cells[2].text = "oui" if a["inondation_documentee"] else ""

    fiab = _fiabilite()
    if fiab:
        doc.add_heading("Fiabilité du système", 1)
        doc.add_paragraph(fiab)
    for a in b.get("avertissements") or []:
        doc.add_paragraph(a, style="List Bullet")

    doc.add_heading("Définitions et limites", 1)
    doc.add_paragraph(
        "Année extrême : saison dont l'empreinte (somme de l'étendue de tous les événements "
        "de pluie extrême CHIRPS) est dans le tiers supérieur. Cette définition retrouve les "
        "inondations majeures de 1999, 2003, 2005, 2009, 2010, 2012, 2020 et 2022. Une "
        "probabilité n'est pas une certitude ; un risque faible n'exclut pas des pluies "
        "intenses locales. Ce bulletin ne remplace pas les prévisions et alertes officielles "
        "de l'ANACIM.")
    doc.save(str(chemin))


def apercu(b):
    """Ce qui sera ecrit, montre avant approbation."""
    return {"sms": sms(b), "sms_caracteres": len(sms(b)), "resume_markdown": resume(b)}


def ecrire(annee):
    """Ecrit les trois versions (appele APRES approbation)."""
    from . import production
    b = production.lire_bulletin(annee)
    if b is None:
        raise FileNotFoundError("Bulletin %d introuvable." % annee)
    DOSSIER.mkdir(parents=True, exist_ok=True)
    base = DOSSIER / ("veille_%d" % annee)
    fichiers = {"sms": base.with_suffix(".sms.txt"), "resume": base.with_suffix(".resume.md"),
                "docx": base.with_suffix(".docx")}
    fichiers["sms"].write_text(sms(b) + "\n", encoding="utf-8")
    fichiers["resume"].write_text(resume(b), encoding="utf-8")
    docx(b, fichiers["docx"])
    return {"fichier": str(fichiers["docx"].relative_to(DOSSIER_SORTIE.parent.parent)),
            "fichiers": [str(f.relative_to(DOSSIER_SORTIE.parent.parent)) for f in fichiers.values()]}
